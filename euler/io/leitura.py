"""Leitura genérica das tabelas do contrato (T03, T05).

Princípios (AGENTS.md, regras 2 e 3):
- o arquivo original é preservado como texto (`Importacao.original`);
- vazio é **ausente**, nunca zero; nada é preenchido ou interpolado;
- todo valor que não pôde ser lido, ou que foi interpretado (vírgula decimal,
  fuso assumido, data dia/mês/ano), gera um aviso com linha e motivo.

Aceita CSV separado por vírgula ou por ponto-e-vírgula (padrão do Excel em
português, com vírgula decimal), em UTF-8 ou na codificação do Windows, e abas
de uma planilha .xlsx já lidas como texto.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Literal

import pandas as pd

from euler.io.esquemas import TABELAS, Coluna, Tabela, rotulo_categoria

FUSO_PADRAO = "America/Sao_Paulo"
Gravidade = Literal["erro", "atencao", "info"]
Fonte = str | Path | bytes | BinaryIO | pd.DataFrame

_VERDADEIRO = {"true", "verdadeiro", "sim", "s", "1", "x", "yes"}
_FALSO = {"false", "falso", "nao", "n", "0", "no"}
_DATA_BR = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?")
_MILHAR_BR = re.compile(r"-?\d{1,3}(\.\d{3})+(,\d+)?")
_MARCAS_SEM_DADO = {"-", "--", "nan", "na", "n/a", "s/d", "sd", "null", "none"}

# Integração adaptativa: aliases só quando a semântica/unidade está explícita.
# Nomes ambíguos como "steam_flow" não são convertidos silenciosamente.
ALIASES_COLUNAS = {
    "timestamp": "instante_observado",
    "datetime": "instante_observado",
    "boiler_id": "caldeira_id",
    "steam_flow_t_h": "vazao_vapor_t_h",
    "steam_flow_tph": "vazao_vapor_t_h",
    "fuel_flow_kg_h": "vazao_combustivel_kg_h",
    "fuel_pci_mj_kg": "pci_combustivel_mj_kg",
    "fuel_thermal_power_mw": "potencia_combustivel_mw",
    "steam_pressure_bar_g": "p_vapor_bar_man",
    "steam_temperature_c": "t_vapor_c",
    "steam_quality_frac": "titulo_vapor_frac",
    "feedwater_temperature_c": "t_agua_alim_c",
    "feedwater_flow_t_h": "vazao_agua_alim_t_h",
    "flue_gas_temperature_c": "t_gases_c",
    "stack_temperature_c": "t_gases_c",
    "o2_dry_pct": "o2_seco_pct",
    "flue_gas_o2_dry_pct": "o2_seco_pct",
    "economizer_water_pressure_bar_g": "p_agua_eco_bar_man",
    "economizer_water_inlet_c": "t_agua_eco_entrada_c",
    "economizer_water_outlet_c": "t_agua_eco_saida_c",
    "economizer_gas_inlet_c": "t_gases_eco_entrada_c",
    "economizer_gas_outlet_c": "t_gases_eco_saida_c",
}


@dataclass(frozen=True)
class Aviso:
    """Um problema ou interpretação encontrado na importação.

    linha: número da linha no arquivo (cabeçalho = 1); None quando vale para a tabela.
    gravidade: "erro" impede usar a tabela; "atencao" afeta análises; "info" é registro.
    """

    tabela: str
    linha: int | None
    coluna: str | None
    tipo: str
    mensagem: str
    gravidade: Gravidade = "atencao"


@dataclass
class Importacao:
    """Resultado da importação de uma tabela.

    original: o arquivo como veio (tudo texto), com a coluna `linha`.
    dados: tabela normalizada (tipos e unidades do contrato), com a coluna `linha`.
    """

    tabela: str
    original: pd.DataFrame
    dados: pd.DataFrame
    avisos: list[Aviso] = field(default_factory=list)

    @property
    def bloqueada(self) -> bool:
        return any(a.gravidade == "erro" for a in self.avisos)


# ---------------------------------------------------------------- leitura bruta


def _bytes(fonte: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(fonte, bytes):
        return fonte
    if isinstance(fonte, str | Path):
        return Path(fonte).read_bytes()
    dados = fonte.read()
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    return dados


def ler_csv(
    fonte: str | Path | bytes | BinaryIO, nome_tabela: str
) -> tuple[pd.DataFrame, bool, list[Aviso]]:
    """Lê um CSV como texto. Retorna (tabela bruta com `linha`, usa_virgula_decimal, avisos)."""
    avisos: list[Aviso] = []
    conteudo = _bytes(fonte)
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252")
        avisos.append(
            Aviso(
                nome_tabela,
                None,
                None,
                "codificacao",
                "Arquivo na codificação do Windows (cp1252); acentos convertidos para UTF-8.",
                "info",
            )
        )
    primeira = texto.split("\n", 1)[0]
    separador = ";" if primeira.count(";") > primeira.count(",") else ","
    leitor = csv.reader(io.StringIO(texto), delimiter=separador)
    try:
        cabecalho = [c.strip() for c in next(leitor)]
    except StopIteration:
        avisos.append(Aviso(nome_tabela, None, None, "arquivo_vazio", "Arquivo vazio.", "erro"))
        return pd.DataFrame(columns=["linha"]), False, avisos
    linhas = []
    for valores in leitor:
        if not any(v.strip() for v in valores):
            continue
        if len(valores) > len(cabecalho):
            avisos.append(
                Aviso(
                    nome_tabela,
                    leitor.line_num,
                    None,
                    "colunas_a_mais",
                    f"Linha com {len(valores)} valores para {len(cabecalho)} colunas; "
                    "os valores a mais foram ignorados (estão no original).",
                )
            )
        valores = (valores + [""] * len(cabecalho))[: len(cabecalho)]
        linhas.append({"linha": leitor.line_num, **dict(zip(cabecalho, valores, strict=True))})
    bruto = pd.DataFrame(linhas, columns=["linha", *cabecalho])
    return bruto, separador == ";", avisos


def ler_planilha(fonte: str | Path | bytes | BinaryIO) -> dict[str, pd.DataFrame]:
    """Lê as abas de uma planilha .xlsx como texto, uma por tabela do contrato.

    Retorna {nome_da_tabela: tabela bruta com `linha`} só para as abas com nome de
    tabela (diario, combustivel, amostras, eventos, instrumentos).
    """
    conteudo = _bytes(fonte)
    abas = pd.read_excel(io.BytesIO(conteudo), sheet_name=None, dtype=str, header=0)
    porcentagens = porcentagens_excel(conteudo)
    resultado = {}
    for nome_aba, df in abas.items():
        nome = nome_aba.strip().lower()
        if nome not in TABELAS:
            continue
        df = df.fillna("").astype(str)
        # cabeçalho na linha 1 do Excel: a linha r da planilha é a linha r - 2 da tabela
        for r, c, texto in porcentagens.get(nome_aba, []):
            if r >= 2 and r - 2 < len(df) and c - 1 < len(df.columns):
                df.iat[r - 2, c - 1] = texto
        df.columns = [str(c).strip() for c in df.columns]
        df = df[[c for c in df.columns if not c.startswith("Unnamed")]]
        if not df.empty:
            df = df[df.apply(lambda r: any(v.strip() for v in r), axis=1)]
        df.insert(0, "linha", df.index + 2)
        resultado[nome] = df.reset_index(drop=True)
    return resultado


def porcentagens_excel(conteudo: bytes) -> dict[str, list[tuple[int, int, str]]]:
    """Células numéricas com formato de porcentagem no Excel: {aba: [(linha, coluna, texto)]}.

    O Excel guarda "45%" como 0,45 e mostra 45%. O texto devolvido é o que a pessoa vê
    ("45%"); a conversão para a unidade da coluna é feita depois por `numero_na_unidade`,
    uma vez só. Linha e coluna começam em 1 (A1 = (1, 1)).
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 — quem lê a planilha já informa o erro de formato
        return {}
    saida: dict[str, list[tuple[int, int, str]]] = {}
    try:
        for ws in wb.worksheets:
            for linha in ws.iter_rows():
                for cel in linha:
                    v = getattr(cel, "value", None)
                    formato = getattr(cel, "number_format", "") or ""
                    if (
                        isinstance(v, int | float)
                        and not isinstance(v, bool)
                        and "%" in formato
                        and cel.row is not None
                    ):
                        saida.setdefault(ws.title, []).append(
                            (cel.row, cel.column, f"{v * 100:.12g}%")
                        )
    finally:
        wb.close()
    return saida


# ---------------------------------------------------------------- conversões

_NOME_TIPO = {
    "numero": "número",
    "instante": "data e hora",
    "data": "data",
    "booleano": "verdadeiro/falso",
    "categoria": "categoria",
    "texto": "texto",
}


def descrever_linhas(numeros, maximo: int = 5) -> str:
    """'linha 4', 'linhas 4 e 5', 'linhas 4, 5, 6 e mais 10'."""
    ns = [str(int(n)) for n in numeros]
    if len(ns) == 1:
        return f"linha {ns[0]}"
    if len(ns) > maximo:
        return f"linhas {', '.join(ns[:maximo])} e mais {len(ns) - maximo}"
    return f"linhas {', '.join(ns[:-1])} e {ns[-1]}"


def _dica_faixa(col: Coluna, abaixo: bool) -> str:
    if isinstance(col.dica_faixa, tuple):
        return col.dica_faixa[0] if abaixo else col.dica_faixa[1]
    if col.dica_faixa:
        return col.dica_faixa
    lo, hi = col.faixa
    return f"Fora da faixa plausível ({lo:g} a {hi:g} {col.unidade})."


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def _numero(texto: str, virgula_decimal: bool) -> tuple[float | None, str | None]:
    """Converte texto em número. Retorna (valor, interpretação feita ou None)."""
    t = texto.strip().replace(" ", "").replace(" ", "")
    interpretacao = None
    if _MILHAR_BR.fullmatch(t) and (virgula_decimal or "," in t):
        t = t.replace(".", "").replace(",", ".")
        interpretacao = "ponto lido como separador de milhar e vírgula como decimal"
    elif "," in t:
        t = t.replace(",", ".")
        interpretacao = "vírgula lida como decimal"
    return float(t), interpretacao


def numero_na_unidade(texto: str, virgula_decimal: bool, unidade: str) -> tuple[float, str | None]:
    """Número do texto expresso em `unidade`. Retorna (valor, interpretação feita ou None).

    Texto com "%" no fim (célula de porcentagem do Excel, ou escrita assim) vale n/100:
    entra como n numa unidade em %, como n/100 numa unidade em fração, e é recusado em
    qualquer outra unidade (ValueError). Sem "%", o número vale como está.
    """
    t = texto.strip()
    if not t.endswith("%"):
        return _numero(t, virgula_decimal)
    valor, _ = _numero(t[:-1], virgula_decimal)
    u = _sem_acento(unidade.strip().lower())
    if u.startswith("%"):
        return valor, "valor com % lido em porcentagem"
    if u.startswith("fracao"):
        return valor / 100, "valor com % convertido para fração (45% → 0,45)"
    raise ValueError(f"'{t}' é porcentagem e a unidade é {unidade}")


def _instante(texto: str, fuso: str) -> tuple[pd.Timestamp, list[str]]:
    """Converte texto em instante com fuso. Retorna (instante, interpretações feitas)."""
    t = texto.strip()
    interpretacoes = []
    m = _DATA_BR.fullmatch(t)
    if m:
        dia, mes, ano, hora, minuto, segundo = (int(g) if g else 0 for g in m.groups())
        ts = pd.Timestamp(year=ano, month=mes, day=dia, hour=hora, minute=minuto, second=segundo)
        interpretacoes.append("data lida como dia/mês/ano")
    else:
        ts = pd.Timestamp(t)
        if pd.isna(ts):
            raise ValueError(t)
    if ts.tzinfo is None:
        ts = ts.tz_localize(fuso)
        interpretacoes.append(
            "sem fuso horário; assumido o horário de Brasília"
            if fuso == FUSO_PADRAO
            else f"sem fuso horário; assumido {fuso}"
        )
    return ts.tz_convert(fuso), interpretacoes


def _data(texto: str) -> tuple[pd.Timestamp, list[str]]:
    t = texto.strip()
    m = _DATA_BR.fullmatch(t)
    if m:
        dia, mes, ano = (int(g) for g in m.groups()[:3])
        return pd.Timestamp(year=ano, month=mes, day=dia), ["data lida como dia/mês/ano"]
    ts = pd.Timestamp(t)
    if pd.isna(ts):
        raise ValueError(t)
    return ts.tz_localize(None).normalize() if ts.tzinfo else ts.normalize(), []


def _converter_coluna(
    tabela: Tabela,
    col: Coluna,
    textos: pd.Series,
    linhas: pd.Series,
    virgula_decimal: bool,
    fuso: str,
) -> tuple[pd.Series, list[Aviso]]:
    avisos: list[Aviso] = []
    interpretadas: dict[str, list[int]] = {}

    def aviso(linha, tipo, msg, grav: Gravidade = "atencao"):
        avisos.append(Aviso(tabela.nome, int(linha), col.nome, tipo, msg, grav))

    valores = []
    for texto, linha in zip(textos, linhas, strict=True):
        texto = str(texto).strip()
        if texto.lower() in _MARCAS_SEM_DADO:
            interpretadas.setdefault(f"'{texto}' lido como ausente", []).append(int(linha))
            texto = ""
        if texto == "":
            if col.obrigatoria:
                aviso(
                    linha, "obrigatorio_vazio", f"Valor obrigatório de “{col.rotulo}” está vazio."
                )
            valores.append(None)
            continue
        try:
            if col.tipo == "numero":
                valor, interp = numero_na_unidade(texto, virgula_decimal, col.unidade)
                if interp:
                    interpretadas.setdefault(interp, []).append(int(linha))
                if col.faixa and not col.faixa[0] <= valor <= col.faixa[1]:
                    dica = _dica_faixa(col, abaixo=valor < col.faixa[0])
                    aviso(
                        linha,
                        "unidade_suspeita",
                        f"{col.rotulo_inicial} = {texto}. {dica} O valor foi mantido.",
                    )
            elif col.tipo == "instante":
                valor, interps = _instante(texto, fuso)
                for interp in interps:
                    interpretadas.setdefault(interp, []).append(int(linha))
            elif col.tipo == "data":
                valor, interps = _data(texto)
                for interp in interps:
                    interpretadas.setdefault(interp, []).append(int(linha))
            elif col.tipo == "booleano":
                chave = _sem_acento(texto.lower())
                if chave in _VERDADEIRO:
                    valor = True
                elif chave in _FALSO:
                    valor = False
                else:
                    raise ValueError(texto)
            elif col.tipo == "categoria":
                valor = _sem_acento(texto.lower()).replace(" ", "_")
                if valor != texto:
                    interpretadas.setdefault(
                        "categoria padronizada (minúsculas, sem acento)", []
                    ).append(int(linha))
                if valor not in col.categorias:
                    aviso(
                        linha,
                        "categoria_desconhecida",
                        f"{col.rotulo_inicial} = '{texto}' não está na lista "
                        f"({', '.join(rotulo_categoria(c) for c in col.categorias)}).",
                    )
            else:
                valor = texto
        except (ValueError, TypeError, OverflowError):
            aviso(
                linha,
                "valor_ilegivel",
                f"{col.rotulo_inicial} = '{texto}' não pôde ser lido como "
                f"{_NOME_TIPO[col.tipo]}; "
                "tratado como ausente.",
            )
            valor = None
        valores.append(valor)

    for interp, linhas_interp in interpretadas.items():
        avisos.append(
            Aviso(
                tabela.nome,
                linhas_interp[0],
                col.nome,
                "interpretacao",
                f"{col.rotulo_inicial}: {interp} ({descrever_linhas(linhas_interp)}).",
                "info",
            )
        )

    if col.tipo == "numero":
        serie = pd.Series(valores, dtype="Float64")
    elif col.tipo == "instante":
        serie = pd.Series(pd.to_datetime(valores, utc=True)).dt.tz_convert(fuso)
    elif col.tipo == "data":
        serie = pd.Series(pd.to_datetime(valores))
    elif col.tipo == "booleano":
        serie = pd.Series(valores, dtype="boolean")
    else:
        serie = pd.Series(valores, dtype="string")
    return serie, avisos


def normalizar(
    nome_tabela: str,
    bruto: pd.DataFrame,
    virgula_decimal: bool = False,
    fuso: str = FUSO_PADRAO,
) -> Importacao:
    """Converte uma tabela bruta (texto, com `linha`) para os tipos do contrato."""
    tabela = TABELAS[nome_tabela]
    avisos: list[Aviso] = []
    original = bruto.reset_index(drop=True)
    entrada = original.copy()
    nomes_contrato = [c.nome for c in tabela.colunas]
    nomes_contrato_set = set(nomes_contrato)

    # Dois nomes para a mesma grandeza exigem resolução explícita, não prioridade silenciosa.
    grupos: dict[str, list[str]] = {}
    for coluna in entrada.columns:
        canonico = ALIASES_COLUNAS.get(coluna, coluna)
        if canonico in nomes_contrato_set:
            grupos.setdefault(canonico, []).append(coluna)
    for canonico, colunas in grupos.items():
        if len(colunas) > 1:
            avisos.append(
                Aviso(
                    tabela.nome,
                    None,
                    canonico,
                    "aliases_ambiguos",
                    f"Mais de uma coluna representa {canonico}: {', '.join(colunas)}. "
                    "Confirme qual usar e envie uma única coluna; o original foi preservado.",
                    "erro",
                )
            )
    if avisos:
        vazio = pd.DataFrame(columns=["linha", *nomes_contrato])
        return Importacao(tabela.nome, original, vazio, avisos)

    renomear = {}
    for alias, canonico in ALIASES_COLUNAS.items():
        if alias not in entrada.columns or canonico in entrada.columns:
            continue
        if canonico not in nomes_contrato_set:
            continue
        renomear[alias] = canonico
        avisos.append(
            Aviso(
                tabela.nome,
                None,
                alias,
                "alias_coluna",
                f"Coluna “{alias}” reconhecida como “{canonico}”; o arquivo original foi preservado.",
                "info",
            )
        )
    if renomear:
        entrada = entrada.rename(columns=renomear)

    presentes = [c for c in entrada.columns if c != "linha"]

    for extra in [c for c in presentes if c not in nomes_contrato]:
        avisos.append(
            Aviso(
                tabela.nome,
                None,
                extra,
                "coluna_desconhecida",
                f"Coluna “{extra}” não faz parte do modelo; ignorada (está no original).",
                "info",
            )
        )
    for col in tabela.colunas:
        if col.nome in presentes:
            continue
        if col.obrigatoria:
            avisos.append(
                Aviso(
                    tabela.nome,
                    None,
                    col.nome,
                    "coluna_obrigatoria_ausente",
                    f"Falta a coluna obrigatória “{col.rotulo}” ({col.descricao}).",
                    "erro",
                )
            )
        else:
            avisos.append(
                Aviso(
                    tabela.nome,
                    None,
                    col.nome,
                    "coluna_ausente",
                    f"Coluna opcional “{col.rotulo}” ausente; tratada como vazia.",
                    "info",
                )
            )

    if any(a.gravidade == "erro" for a in avisos):
        vazio = pd.DataFrame(columns=["linha", *nomes_contrato])
        return Importacao(tabela.nome, original, vazio, avisos)

    dados = pd.DataFrame({"linha": original["linha"].astype(int)})
    for col in tabela.colunas:
        textos = entrada[col.nome] if col.nome in presentes else pd.Series([""] * len(original))
        serie, avisos_col = _converter_coluna(
            tabela, col, textos, original["linha"], virgula_decimal, fuso
        )
        dados[col.nome] = serie.array
        avisos.extend(avisos_col)
    return Importacao(tabela.nome, original, dados, avisos)


def importar_tabela(nome_tabela: str, fonte: Fonte, fuso: str = FUSO_PADRAO) -> Importacao:
    """Lê (CSV ou aba já lida) e normaliza uma tabela do contrato, sem regras específicas."""
    if isinstance(fonte, pd.DataFrame):
        return normalizar(nome_tabela, fonte, virgula_decimal=False, fuso=fuso)
    bruto, virgula, avisos = ler_csv(fonte, nome_tabela)
    if any(a.gravidade == "erro" for a in avisos):
        return Importacao(nome_tabela, bruto, pd.DataFrame(columns=["linha"]), avisos)
    imp = normalizar(nome_tabela, bruto, virgula_decimal=virgula, fuso=fuso)
    imp.avisos[:0] = avisos
    return imp
