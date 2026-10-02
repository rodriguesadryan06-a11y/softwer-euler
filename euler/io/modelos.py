"""Gera o contrato de dados em Markdown e a planilha modelo .xlsx a partir dos esquemas (T02).

Rodar: python scripts/gerar_modelos.py
"""

from __future__ import annotations

import csv
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from euler.formato import num
from euler.io.esquemas import TABELAS, Coluna

RAIZ = Path(__file__).resolve().parents[2]
PASTA_MODELOS = RAIZ / "templates"
ARQUIVO_CONTRATO = RAIZ / "docs" / "contrato_dados.md"
ARQUIVO_PLANILHA = PASTA_MODELOS / "planilha_modelo_euler.xlsx"

REGRAS_GERAIS = """\
- **Formato recomendado:** um arquivo por tabela (`diario.csv`, `combustivel.csv`,
  `amostras.csv`, `eventos.csv`, `instrumentos.csv`) ou uma planilha `.xlsx` com uma aba
  por tabela. Um CSV com outro nome também pode ser reconhecido pelo cabeçalho quando a tabela
  for inequívoca; aliases só são aceitos quando a unidade está explícita no nome.
- **Separador:** vírgula (`,`) ou ponto-e-vírgula (`;`, padrão do Excel em português). Com
  ponto-e-vírgula, números podem usar vírgula decimal (`182,5`). Codificação UTF-8 ou Windows.
- **Vazio = não medido.** Nunca escreva 0 para "não medido". `-`, `s/d` e `n/a` também são
  lidos como vazio (com aviso).
- **Datas e horas:** ISO 8601 com fuso (`2026-10-05T08:00:00-03:00`). Também é aceito
  `05/10/2026 08:00` (dia/mês/ano); sem fuso, assume-se o horário de Brasília, com aviso.
- **Unidade no nome da coluna.** A pressão do diário é **manométrica** (`p_vapor_bar_man`); a
  EULER converte para absoluta com a pressão atmosférica do local (altitude ou barômetro).
- **Umidade em fração, base úmida** (`0,38` = 38%). Composição em fração, base seca.
- **`origem_dado`:** `sintetico`, `publico` ou `real`. Dados reais de clientes nunca entram no
  repositório do código.
- **O que o importador faz:** guarda o arquivo original; converte unidades; lista avisos com
  linha e motivo (lacunas, duplicatas, totalizador reiniciado, registro tardio, unidades
  suspeitas, volume sem densidade). **Nada é preenchido, interpolado ou corrigido em silêncio.**
"""


def _faixa(col: Coluna) -> str:
    if not col.faixa:
        return ""
    lo, hi = col.faixa
    return f"{num(lo, 0 if lo == int(lo) else 2)} a {num(hi, 0 if hi == int(hi) else 2)}"


def _descricao(col: Coluna) -> str:
    if col.categorias:
        return f"{col.descricao} Valores: {', '.join(f'`{c}`' for c in col.categorias)}."
    return col.descricao


def contrato_markdown() -> str:
    """Texto completo de docs/contrato_dados.md."""
    partes = [
        "# Contrato de dados · EULER",
        "",
        (
            "> Gerado automaticamente a partir de `euler/io/esquemas.py` por "
            "`python scripts/gerar_modelos.py`. **Não edite à mão**: mude o esquema e gere de novo."
        ),
        (
            "> Derivado dos modelos de CSV do kit de construção; a spec v0.3 não estava "
            "disponível (ver D12 em `docs/decisoes.md`)."
        ),
        "",
        "## Regras gerais",
        "",
        REGRAS_GERAIS,
        "## Tabelas",
    ]
    for t in TABELAS.values():
        partes += [
            "",
            f"### `{t.arquivo}` · {t.titulo}",
            "",
            t.descricao,
            "",
            "| Coluna | Unidade | Obrigatória | Descrição | Exemplo | Faixa plausível (só aviso) |",
            "|---|---|---|---|---|---|",
        ]
        for c in t.colunas:
            exemplo = f"`{c.exemplo}`" if c.exemplo else ""
            partes.append(
                f"| `{c.nome}` | {c.unidade} | {'sim' if c.obrigatoria else 'não'} | "
                f"{_descricao(c)} | {exemplo} | {_faixa(c)} |"
            )
    return "\n".join(partes) + "\n"


INSTRUCOES_PLANILHA = [
    "EULER · Planilha modelo de registros da caldeira",
    "",
    "Como preencher",
    (
        "1. Use uma aba para cada tipo de registro: diario, combustivel, amostras, eventos, "
        "instrumentos. Não mude os nomes das abas nem das colunas."
    ),
    "2. Uma linha por leitura, recebimento, amostra ou evento.",
    "3. Deixe em branco o que não foi medido. Nunca escreva 0 para 'não medido'.",
    "4. Datas e horas: 05/10/2026 08:00 (horário de Brasília) ou 2026-10-05T08:00:00-03:00.",
    "5. Números podem usar vírgula decimal (182,5).",
    "6. A unidade está no nome da coluna. Passe o mouse no cabeçalho para ver a explicação.",
    "7. Umidade em fração: 0,38 para 38%.",
    "8. A linha 2 de cada aba é um exemplo SINTÉTICO: apague antes de usar com dados reais.",
    "",
    "Como enviar",
    (
        "Salve esta planilha (.xlsx) e envie na tela 'Importar dados' da EULER. Também é "
        "possível exportar cada aba como CSV com o nome da aba (diario.csv, combustivel.csv...)."
    ),
    "",
    (
        "A EULER guarda o arquivo original e mostra uma lista de avisos (linha e motivo). "
        "Ela nunca preenche ou corrige valores sem avisar."
    ),
]


def _exemplo_do_modelo_csv(nome: str) -> list[str]:
    with (PASTA_MODELOS / f"{nome}.csv").open(encoding="utf-8") as f:
        linhas = list(csv.reader(f))
    return linhas[1] if len(linhas) > 1 else []


def gerar_planilha(destino: Path = ARQUIVO_PLANILHA) -> Path:
    """Cria a planilha modelo com instruções, uma aba por tabela e a linha de exemplo."""
    wb = Workbook()
    leia = wb.active
    leia.title = "LEIA-ME"
    for i, texto in enumerate(INSTRUCOES_PLANILHA, start=1):
        cel = leia.cell(row=i, column=1, value=texto)
        if i == 1:
            cel.font = Font(bold=True, size=14)
        elif texto in ("Como preencher", "Como enviar"):
            cel.font = Font(bold=True)
    leia.column_dimensions["A"].width = 110

    cabecalho = PatternFill("solid", fgColor="FDE7D9")
    for t in TABELAS.values():
        aba = wb.create_sheet(t.nome)
        exemplo = _exemplo_do_modelo_csv(t.nome)
        for j, col in enumerate(t.colunas, start=1):
            cel = aba.cell(row=1, column=j, value=col.nome)
            cel.font = Font(bold=True)
            cel.fill = cabecalho
            cel.alignment = Alignment(horizontal="center")
            obrig = "Obrigatória" if col.obrigatoria else "Opcional"
            cel.comment = Comment(
                f"{_descricao(col)}\nUnidade: {col.unidade}\n{obrig}",
                "EULER",
                width=320,
                height=120,
            )
            letra = cel.column_letter
            aba.column_dimensions[letra].width = max(14, len(col.nome) + 4)
            valor = exemplo[j - 1] if j - 1 < len(exemplo) else ""
            if col.tipo == "numero":
                if valor:
                    aba.cell(row=2, column=j, value=float(valor) if "." in valor else int(valor))
            else:
                alvo = aba.cell(row=2, column=j, value=valor or None)
                alvo.number_format = "@"
                for linha in range(3, 501):
                    aba.cell(row=linha, column=j).number_format = "@"
            if col.categorias:
                validacao = DataValidation(
                    type="list", formula1=f'"{",".join(col.categorias)}"', allow_blank=True
                )
                validacao.error = f"Use um destes valores: {', '.join(col.categorias)}"
                aba.add_data_validation(validacao)
                validacao.add(f"{letra}2:{letra}500")
        aba.freeze_panes = "A2"

    wb.properties.creator = "EULER"
    # data fixa para o arquivo gerado não mudar a cada geração
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1, tzinfo=UTC)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def gerar_tudo() -> list[Path]:
    ARQUIVO_CONTRATO.write_text(contrato_markdown(), encoding="utf-8")
    return [ARQUIVO_CONTRATO, gerar_planilha()]
