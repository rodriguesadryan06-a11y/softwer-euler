"""Contrato de dados e modelos sempre em dia com os esquemas (T02)."""

import csv

from openpyxl import load_workbook

from euler.io.esquemas import TABELAS
from euler.io.modelos import ARQUIVO_CONTRATO, ARQUIVO_PLANILHA, PASTA_MODELOS, contrato_markdown


def test_contrato_em_dia_com_os_esquemas():
    atual = ARQUIVO_CONTRATO.read_text(encoding="utf-8")
    assert atual == contrato_markdown(), "rode: python scripts/gerar_modelos.py"


def test_toda_coluna_tem_unidade_obrigatoriedade_e_descricao():
    for t in TABELAS.values():
        for c in t.colunas:
            assert c.unidade and c.descricao, f"{t.nome}.{c.nome}"


def test_cabecalhos_dos_csvs_modelo_seguem_o_contrato():
    """Os modelos do kit vêm na ordem do contrato; colunas acrescentadas depois ao contrato
    (ex.: incerteza_tipo, Fase R) ficam no fim e são opcionais."""
    for t in TABELAS.values():
        with (PASTA_MODELOS / t.arquivo).open(encoding="utf-8") as f:
            cabecalho = next(csv.reader(f))
        contrato = [c.nome for c in t.colunas]
        assert cabecalho == contrato[: len(cabecalho)], t.arquivo
        assert all(not t.coluna(c).obrigatoria for c in contrato[len(cabecalho) :]), t.arquivo


def test_planilha_tem_uma_aba_por_tabela_com_cabecalhos_compativeis_com_o_contrato():
    """A planilha versionada pode ficar sem colunas opcionais recém-adicionadas.

    O importador normaliza essas colunas como vazias. Quando o modelo for regenerado pelo
    script, ele passa a conter todas; nenhuma coluna obrigatória pode faltar.
    """
    wb = load_workbook(ARQUIVO_PLANILHA)
    assert wb.sheetnames == ["LEIA-ME", *TABELAS]
    for t in TABELAS.values():
        cabecalho = [c.value for c in wb[t.nome][1]]
        contrato = [c.nome for c in t.colunas]
        assert cabecalho == contrato[: len(cabecalho)], t.nome
        assert all(not t.coluna(c).obrigatoria for c in contrato[len(cabecalho) :]), t.nome
