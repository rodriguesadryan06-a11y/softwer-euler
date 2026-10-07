"""Planilha de fábrica como ela chega (D106): vocabulário de fábrica, unidade no cabeçalho,
título acima do cabeçalho e Data + Hora em colunas separadas. Dados sintéticos."""

import io
import json

import pandas as pd
import pytest
from openpyxl import Workbook

from app.importacao_guiada import (
    combinar_data_hora,
    ler_fontes,
    preparar_lote,
    resumo_da_fonte,
    sugerir_combinacao,
    sugerir_mapeamento,
    sugestoes_detalhadas,
    unidades_permitidas,
)
from euler.io import fontes_de_arquivos, importar_pacote
from euler.io.esquemas import TABELAS

CAB_DIARIO = [
    "Data",
    "Hora",
    "Pressão vapor (kgf/cm²)",
    "Temp. chaminé (°C)",
    "O2 base seca (%)",
    "Totalizador vapor (t)",
    "Observações",
    "Assinatura",
]


def planilha_da_fabrica() -> bytes:
    """Duas abas como uma fábrica costuma mandar: título em cima, Data e Hora separadas."""
    wb = Workbook()
    d = wb.active
    d.title = "Diário caldeira"
    d.append(["Indústria Exemplo (sintética) · Diário da caldeira 1"])
    d.append(["Setembro/2026"])
    d.append([])
    d.append(CAB_DIARIO)
    d.append(["01/09/2026", "08:00", "8,5", "182", "7,9", "15200", "ok", "J."])
    d.append(["01/09/2026", "10:00", "8,7", "", "8,1", "15212", "", "J."])
    d.append(["01/09/2026", "", "8,6", "185", "8,0", "15224", "sem hora", "M."])
    lenha = wb.create_sheet("Entrada de lenha")
    lenha.append(["Data", "Fornecedor", "Peso líquido (t)", "Valor total (R$)", "Nota fiscal"])
    lenha.append(["01/09/2026", "F1", "28,4", "4260", "1234"])
    lenha.append(["03/09/2026", "F2", "30,1", "4515", "1250"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------ vocabulário de fábrica


def test_vocabulario_de_fabrica_sugere_com_motivo_e_unidade():
    s = sugestoes_detalhadas(CAB_DIARIO, "diario")
    assert s["Pressão vapor (kgf/cm²)"].alvo == "p_vapor_bar_man"
    assert s["Pressão vapor (kgf/cm²)"].unidade == "kgf/cm² manométrico"
    assert s["Temp. chaminé (°C)"].alvo == "t_gases_c"
    assert s["Temp. chaminé (°C)"].unidade == "°C"
    assert s["O2 base seca (%)"].alvo == "o2_seco_pct"
    assert s["Totalizador vapor (t)"].alvo == "totalizador_vapor_t"
    assert s["Totalizador vapor (t)"].unidade == "t (acumulado)"
    assert all(x.motivo for x in s.values())
    # sem regra para observações e assinatura: ficam de fora, guardadas no original
    assert "Observações" not in s and "Assinatura" not in s


def test_cabecalho_ambiguo_continua_sem_sugestao():
    mapa = sugerir_mapeamento(["Pressão", "Consumo", "R$", "Temperatura", "Peso"], "diario")
    assert mapa == {}
    # O₂ sem dizer a base: a EULER usa base seca e não presume
    assert "O2 (%)" not in sugerir_mapeamento(["O2 (%)"], "diario")
    # pressão absoluta não vira manométrica
    assert "Pressão vapor (bar abs)" not in sugerir_mapeamento(
        ["Pressão vapor (bar abs)"], "diario"
    )


def test_valor_por_tonelada_e_peso_bruto_nao_viram_total_nem_massa():
    colunas = ["Data", "Peso bruto (t)", "Tara (t)", "Preço (R$/t)", "Valor total (R$)"]
    mapa = sugerir_mapeamento(colunas, "combustivel")
    assert mapa == {"Data": "data", "Valor total (R$)": "preco_brl"}


def test_duas_colunas_candidatas_ao_mesmo_campo_ficam_para_a_pessoa():
    mapa = sugerir_mapeamento(["Temp. chaminé (°C)", "Temperatura gases saída (°C)"], "diario")
    assert "t_gases_c" not in mapa.values()


def test_mapeamento_salvo_e_nome_exato_tem_prioridade():
    s = sugestoes_detalhadas(
        ["Temp. chaminé (°C)", "t_ar_c"], "diario", {"Temp. chaminé (°C)": "t_ar_c"}
    )
    assert s["Temp. chaminé (°C)"].alvo == "t_ar_c"
    assert "salv" in s["Temp. chaminé (°C)"].motivo
    assert "t_ar_c" not in s  # o destino já foi escolhido pelo mapeamento salvo


# ------------------------------------------------------------ unidades novas (conversão exata)


@pytest.mark.parametrize(
    ("coluna", "unidade", "entrada", "esperado"),
    [
        ("p_vapor_bar_man", "kgf/cm² manométrico", 10, 9.80665),
        ("p_vapor_bar_man", "psi manométrico", 100, 6.894757293168),
        ("p_vapor_bar_man", "MPa manométrico", 1, 10),
        ("t_gases_c", "°F", 212, 100),
        ("vazao_vapor_t_h", "t/dia", 48, 2),
    ],
)
def test_conversoes_exatas_novas(coluna, unidade, entrada, esperado):
    col = TABELAS["diario"].coluna(coluna)
    fator, offset = unidades_permitidas(col)[unidade]
    assert entrada * fator + offset == pytest.approx(esperado, rel=1e-12)


def test_pressao_absoluta_continua_fora_das_opcoes():
    col = TABELAS["diario"].coluna("p_vapor_bar_man")
    assert not any("abs" in u for u in unidades_permitidas(col))


# ------------------------------------------------------------ cabeçalho fora da primeira linha


def test_titulo_acima_do_cabecalho_e_detectado_e_linhas_seguem_o_excel():
    fontes = {f.aba: f for f in ler_fontes({"caldeira.xlsx": planilha_da_fabrica()})}
    diario = fontes["Diário caldeira"]
    assert diario.linha_cabecalho == 4
    assert [c for c in diario.bruto if c != "linha"] == CAB_DIARIO
    assert diario.bruto["linha"].tolist() == [5, 6, 7]
    assert fontes["Entrada de lenha"].linha_cabecalho == 1


def test_csv_com_titulo_e_ajuste_manual_da_linha_do_cabecalho():
    texto = "Relatório do turno;;\n;;\nData;Hora;Vazão vapor (t/h)\n01/09/2026;08:00;2,1\n"
    arquivos = {"turno.csv": texto.encode()}
    f = ler_fontes(arquivos)[0]
    assert f.linha_cabecalho == 3
    assert f.virgula_decimal
    assert f.bruto["linha"].tolist() == [4]
    # a pessoa pode corrigir a linha do cabeçalho; o original continua o mesmo
    g = ler_fontes(arquivos, {"turno.csv": 1})[0]
    assert g.linha_cabecalho == 1


# ------------------------------------------------------------ Data + Hora separadas


def test_data_e_hora_separadas_sao_sugeridas_e_ausente_continua_ausente():
    f = {x.aba: x for x in ler_fontes({"caldeira.xlsx": planilha_da_fabrica()})}["Diário caldeira"]
    combinacao = sugerir_combinacao(f, "diario", {})
    assert combinacao == {"alvo": "instante_observado", "data": "Data", "hora": "Hora"}
    valores = combinar_data_hora(f.bruto["Data"], f.bruto["Hora"]).tolist()
    assert valores == ["01/09/2026 08:00", "01/09/2026 10:00", ""]


@pytest.mark.parametrize(
    ("data", "hora", "esperado"),
    [
        (
            "2026-09-01 00:00:00",
            "08:00:00",
            "2026-09-01 08:00:00",
        ),  # células de data e hora do Excel
        ("01/09/2026", "8h30", "01/09/2026 08:30"),
        ("01/09/2026", "8h", "01/09/2026 08:00"),
        ("", "08:00", ""),
    ],
)
def test_formatos_comuns_de_data_e_hora(data, hora, esperado):
    assert combinar_data_hora(pd.Series([data]), pd.Series([hora])).tolist() == [esperado]


# ------------------------------------------------------------ ponta a ponta e resumo


def _decisoes_sugeridas(arquivos):
    decisoes = {}
    for f in ler_fontes(arquivos):
        tabela = {"Diário caldeira": "diario", "Entrada de lenha": "combustivel"}[f.aba]
        colunas = [c for c in f.bruto if c != "linha"]
        combinacao = sugerir_combinacao(f, tabela, {})
        usadas = {combinacao["data"], combinacao["hora"]} if combinacao else set()
        s = sugestoes_detalhadas([c for c in colunas if c not in usadas], tabela)
        decisoes[f.chave] = {
            "tabela": tabela,
            "mapeamento": {c: x.alvo for c, x in s.items()},
            "unidades": {c: x.unidade for c, x in s.items() if x.unidade},
            "constantes": {"origem_dado": "sintetico", "caldeira_id": "CALD-1"}
            if tabela == "diario"
            else {"origem_dado": "sintetico", "tipo": "recebimento"},
            "linha_cabecalho": f.linha_cabecalho,
            **({"combinar": combinacao} if combinacao else {}),
        }
    return decisoes


def test_planilha_da_fabrica_entra_com_conversoes_e_rastreio():
    arquivos = {"caldeira.xlsx": planilha_da_fabrica()}
    decisoes = _decisoes_sugeridas(arquivos)
    lote = preparar_lote(arquivos, decisoes)
    p = importar_pacote(fontes_de_arquivos(lote)[0], p_atm_bar=1.0)
    diario = p.dados("diario")
    assert diario is not None and len(diario) == 3
    assert diario["p_vapor_bar_man"].iloc[0] == pytest.approx(8.5 * 0.980665)
    assert diario["instante_observado"].iloc[0].hour == 8
    assert pd.isna(diario["instante_observado"].iloc[2]) or p.importacoes["diario"].bloqueada
    assert pd.isna(diario["t_gases_c"].iloc[1])  # vazio no original continua ausente
    comb = p.dados("combustivel")
    assert comb["massa_kg"].tolist() == pytest.approx([28400, 30100])
    assert comb["preco_brl"].tolist() == [4260, 4515]
    m = json.loads(lote["euler_importacao.json"])
    ad = {a["aba"]: a for a in m["adaptacoes"]}
    assert ad["Diário caldeira"]["linha_cabecalho"] == 4
    assert ad["Diário caldeira"]["combinar"] == {
        "alvo": "instante_observado",
        "data": "Data",
        "hora": "Hora",
    }
    assert ad["Diário caldeira"]["linhas_originais"] == [5, 6, 7]
    assert set(ad["Diário caldeira"]["colunas_nao_usadas"]) == {"Observações", "Assinatura"}


def test_resumo_diz_o_que_entrou_o_que_ficou_de_fora_e_o_que_falta():
    arquivos = {"caldeira.xlsx": planilha_da_fabrica()}
    decisoes = _decisoes_sugeridas(arquivos)
    f = {x.aba: x for x in ler_fontes(arquivos)}["Diário caldeira"]
    r = resumo_da_fonte(f, decisoes[f.chave])
    entram = " ".join(r["entram"])
    assert "pressão do vapor" in entram and "kgf/cm²" in entram
    assert "“Data” + “Hora”" in entram
    assert r["fora"] == ["Observações", "Assinatura"]
    faltam = " ".join(r["faltam"])
    assert "temperatura do ar" in faltam and "perda nos gases" in faltam
    assert "temperatura da água de alimentação" in faltam


# ------------------------------------------------------------ parte 2 (D107)


def planilha_mes_por_aba(cabecalho_duplo: bool = True) -> bytes:
    """Um mês por aba, cabeçalho em duas linhas (células mescladas) e estoque em aba própria."""
    wb = Workbook()
    for i, (aba, dia) in enumerate((("Ago", "31/08/2026"), ("Set", "01/09/2026"))):
        ws = wb.active if i == 0 else wb.create_sheet(aba)
        ws.title = aba
        ws.append([f"Diário da caldeira 1 · {aba}/2026 (sintético)"])
        if cabecalho_duplo:
            ws.append(["Data", "Hora", "Temperaturas", "", "", "Pressão vapor (kgf/cm²)"])
            ws.append(["", "", "Gases (°C)", "Água alimentação (°C)", "Ar (°C)", ""])
            ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=5)
        else:
            ws.append(["Data", "Hora", "Temp. chaminé (°C)", "Pressão vapor (kgf/cm²)"])
        for hora, tg in (("08:00", "181"), ("16:00", "184")):
            linha = (
                [dia, hora, tg, "105", "28", "8,5"] if cabecalho_duplo else [dia, hora, tg, "8,5"]
            )
            ws.append(linha)
    receb = wb.create_sheet("Entrada de lenha")
    receb.append(["Data", "Lote", "Peso líquido (t)"])
    receb.append(["31/08/2026", "L-31", "28,4"])
    estoque = wb.create_sheet("Estoque do pátio")
    estoque.append(["Data", "Lote", "Peso líquido (t)"])
    estoque.append(["31/08/2026", "EST-08", "112"])
    estoque.append(["30/09/2026", "EST-09", "96"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_cabecalho_em_duas_linhas_vira_grupo_e_subcoluna():
    fontes = {f.aba: f for f in ler_fontes({"mensal.xlsx": planilha_mes_por_aba()})}
    ago = fontes["Ago"]
    assert ago.cabecalho_duplo and ago.linha_cabecalho == 2
    assert [c for c in ago.bruto if c != "linha"] == [
        "Data",
        "Hora",
        "Temperaturas · Gases (°C)",
        "Temperaturas · Água alimentação (°C)",
        "Temperaturas · Ar (°C)",
        "Pressão vapor (kgf/cm²)",
    ]
    assert ago.bruto["linha"].tolist() == [4, 5]  # linhas do Excel
    s = sugestoes_detalhadas([c for c in ago.bruto if c != "linha"], "diario")
    assert s["Temperaturas · Gases (°C)"].alvo == "t_gases_c"
    assert s["Temperaturas · Água alimentação (°C)"].alvo == "t_agua_alim_c"
    assert s["Temperaturas · Ar (°C)"].alvo == "t_ar_c"
    # aba de recebimentos com cabeçalho simples continua simples
    assert not fontes["Entrada de lenha"].cabecalho_duplo


def test_cabecalho_simples_com_titulo_nao_vira_duplo():
    fontes = {f.aba: f for f in ler_fontes({"caldeira.xlsx": planilha_da_fabrica()})}
    assert not fontes["Diário caldeira"].cabecalho_duplo
    f = ler_fontes({"m.xlsx": planilha_mes_por_aba(cabecalho_duplo=False)})[0]
    assert not f.cabecalho_duplo and f.linha_cabecalho == 2


def test_pessoa_pode_desligar_o_cabecalho_duplo():
    arquivos = {"mensal.xlsx": planilha_mes_por_aba()}
    f = ler_fontes(arquivos, duplos={"mensal.xlsx::Ago": False})[0]
    assert not f.cabecalho_duplo
    assert "Temperaturas" in f.bruto.columns


def _decisoes_mensais(arquivos):
    decisoes = {}
    for f in ler_fontes(arquivos):
        if f.aba in {"Ago", "Set"}:
            tabela, constantes = "diario", {"origem_dado": "sintetico", "caldeira_id": "CALD-1"}
        else:
            tipo = "estoque" if "Estoque" in f.aba else "recebimento"
            tabela, constantes = "combustivel", {"origem_dado": "sintetico", "tipo": tipo}
        colunas = [c for c in f.bruto if c != "linha"]
        combinacao = sugerir_combinacao(f, tabela, {})
        usadas = {combinacao["data"], combinacao["hora"]} if combinacao else set()
        s = sugestoes_detalhadas([c for c in colunas if c not in usadas], tabela)
        decisoes[f.chave] = {
            "tabela": tabela,
            "mapeamento": {c: x.alvo for c, x in s.items()},
            "unidades": {c: x.unidade for c, x in s.items() if x.unidade},
            "constantes": constantes,
            "linha_cabecalho": f.linha_cabecalho,
            "cabecalho_duplo": f.cabecalho_duplo,
            **({"combinar": combinacao} if combinacao else {}),
        }
    return decisoes


def test_um_mes_por_aba_e_estoque_em_aba_propria_viram_uma_tabela_cada():
    arquivos = {"mensal.xlsx": planilha_mes_por_aba()}
    lote = preparar_lote(arquivos, _decisoes_mensais(arquivos))
    p = importar_pacote(fontes_de_arquivos(lote)[0], p_atm_bar=1.0)
    diario = p.dados("diario")
    assert len(diario) == 4
    assert diario["t_agua_alim_c"].tolist() == [105, 105, 105, 105]
    comb = p.dados("combustivel")
    assert sorted(comb["tipo"]) == ["estoque", "estoque", "recebimento"]
    assert comb.loc[comb["tipo"] == "recebimento", "massa_kg"].tolist() == pytest.approx([28400])
    m = json.loads(lote["euler_importacao.json"])
    ad = {a["aba"]: a for a in m["adaptacoes"]}
    assert ad["Ago"]["linhas_no_adaptado"] == [2, 3]
    assert ad["Set"]["linhas_no_adaptado"] == [4, 5]
    assert ad["Set"]["cabecalho_duplo"] is True
    assert ad["Estoque do pátio"]["linhas_no_adaptado"] == [3, 4]


def test_mesmo_registro_em_duas_abas_com_valores_diferentes_para():
    wb = Workbook()
    for aba, temp in (("Ago", "181"), ("Ago (cópia)", "190")):
        ws = wb.active if aba == "Ago" else wb.create_sheet(aba)
        ws.title = aba
        ws.append(["Data", "Hora", "Temp. chaminé (°C)"])
        ws.append(["31/08/2026", "08:00", temp])
    out = io.BytesIO()
    wb.save(out)
    arquivos = {"dup.xlsx": out.getvalue()}
    decisoes = {
        f.chave: {
            "tabela": "diario",
            "mapeamento": {"Temp. chaminé (°C)": "t_gases_c"},
            "unidades": {"Temp. chaminé (°C)": "°C"},
            "constantes": {"origem_dado": "sintetico", "caldeira_id": "CALD-1"},
            "combinar": {"alvo": "instante_observado", "data": "Data", "hora": "Hora"},
        }
        for f in ler_fontes(arquivos)
    }
    with pytest.raises(ValueError, match="valores diferentes") as erro:
        preparar_lote(arquivos, decisoes)
    assert "Ago (cópia)" in str(erro.value) and "linha 2" in str(erro.value)
